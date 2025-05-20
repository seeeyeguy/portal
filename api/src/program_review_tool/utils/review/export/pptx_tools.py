"""
`Program Review Tool` `Program` review export utils file
containing common functionality for modifying a pptx.Presentation
object.
"""

# pylint: disable=wrong-import-order
from pptx import Presentation
from pptx.slide import Slide
from typing import Optional, Tuple


def delete_slide(presentation: Presentation, slide_index: int) -> None:
    """
    Deletes a slide from a presentation based on its 0-indexed position.

    Accepts:
        * presentation (Presentation): The presentation.
        * slide_index (int): Zero-based slide index.

    Returns:
        * None
    """

    slide_id_list = list(presentation.slides._sldIdLst)

    try:
        slide_id = slide_id_list[slide_index]
    except IndexError:
        raise ValueError(f"Slide index {slide_index} is out of range.")

    relationship_id = slide_id.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    if relationship_id is None:
        raise ValueError(f"Relationship ID not found for slide index {slide_index}.")

    presentation.part.drop_rel(relationship_id)
    presentation.slides._sldIdLst.remove(slide_id)


def find_slide_index_and_slide_by_title(
    presentation: Presentation, target_title: str, start_index: int = 1
) -> Optional[Tuple[int, Slide]]:
    """
    Finds and returns the slide index and slide matching the target title.

    Accepts:
        * presentation (Presentation): The presentation object.
        * target_title (str): The slide title to look for.
        * start_index (int): Start index offset used for slide numbering.

    Returns:
        * Optional[Tuple[int,Slide]]: The slide index and Slide if found;
            otherwise, None.
    """

    for idx, slide in enumerate(presentation.slides, start=start_index):
        if (
            slide.shapes.title
            and slide.shapes.title.text.strip() == target_title.strip()
        ):
            return (idx, slide)
    return None
